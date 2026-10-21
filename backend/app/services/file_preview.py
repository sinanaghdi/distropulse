from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_ROWS = 10_000
MAX_FILE_BYTES = 5 * 1024 * 1024


class PreviewError(ValueError):
    """Safe, user-facing file validation error."""


def _value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _headers(values: list[Any]) -> list[str]:
    result = [str(value).strip() if value is not None else "" for value in values]
    while result and not result[-1]:
        result.pop()
    if not result or not any(result):
        raise PreviewError("The first row must contain column headers.")
    if any(not item for item in result):
        raise PreviewError("A column header is blank; every column must have a name.")
    if len(set(result)) != len(result):
        raise PreviewError("Duplicate column headers were found; each header must be unique.")
    return result


def _dict_rows(headers: list[str], rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    if len(rows) > MAX_ROWS:
        raise PreviewError(f"The file exceeds {MAX_ROWS:,} data rows; split it into smaller files.")
    return [
        {header: _value(row[index] if index < len(row) else None) for index, header in enumerate(headers)}
        for row in rows
    ]


def read_table(filename: str, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Read CSV/XLSX in memory. Nothing is written to disk."""
    if not content:
        raise PreviewError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise PreviewError("File size exceeds the 5 MiB limit.")
    suffix = Path(filename or "").suffix.lower()

    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise PreviewError("CSV encoding must be UTF-8.") from exc
        try:
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            raw = list(csv.reader(io.StringIO(text), dialect))
        except csv.Error as exc:
            raise PreviewError("The CSV structure could not be read.") from exc
        if not raw:
            raise PreviewError("The CSV contains no rows.")
        columns = _headers(raw[0])
        rows = [tuple(row) for row in raw[1:] if any(str(cell).strip() for cell in row)]
        return columns, _dict_rows(columns, rows)

    if suffix == ".xlsx":
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise PreviewError("The Excel workbook is invalid or corrupted.")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                iterator = workbook.active.iter_rows(values_only=True)
                first = next(iterator, None)
                if first is None:
                    raise PreviewError("The Excel worksheet is empty.")
                columns = _headers(list(first))
                rows = [tuple(row) for row in iterator if any(cell is not None and str(cell).strip() for cell in row)]
                return columns, _dict_rows(columns, rows)
            finally:
                workbook.close()
        except PreviewError:
            raise
        except Exception as exc:
            raise PreviewError("Could not read the Excel workbook. Save it again and retry.") from exc

    raise PreviewError("Unsupported format. Choose a CSV or XLSX file.")
